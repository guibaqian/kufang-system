"""对比：下载的单表 xlsx vs 原工作簿中的同名表（值/公式、样式、合并、列宽、行高、冻结）+ LibreOffice 重算值 vs 文件缓存值"""
import openpyxl, sys, subprocess, os, tempfile, shutil, json
def style_sig(c):
    f=c.font; b=c.border; fl=c.fill; a=c.alignment
    return (f.name,f.sz,f.b,f.i,str(f.color.rgb if f.color else None),fl.fill_type,str(fl.fgColor.rgb if fl.fgColor else None),
            b.left.style,b.right.style,b.top.style,b.bottom.style,c.number_format,a.horizontal,a.vertical,a.wrap_text)
def compare(orig_path, sheet, out_path, row_map=None, ignore_rows=(), max_row=None):
    """max_row：只比到这一行（下载去掉了后面多余的空行时用）"""
    wo=openpyxl.load_workbook(orig_path)[sheet]; wn_book=openpyxl.load_workbook(out_path); 
    res={'sheetnames':wn_book.sheetnames}
    wn=wn_book[sheet]
    diffs_val=[]; diffs_sty=[]; n=0
    maxr=max_row or max(wo.max_row,wn.max_row); maxc=max(wo.max_column,wn.max_column)
    for r in range(1,maxr+1):
        if r in ignore_rows: continue
        for c in range(1,maxc+1):
            a=wo.cell(r,c); b=wn.cell(r,c); n+=1
            va=a.value; vb=b.value
            if isinstance(va,str) and va.startswith('='): va=va.replace(' ','')
            if isinstance(vb,str) and vb.startswith('='): vb=vb.replace(' ','')
            if va!=vb and not (va in (None,'') and vb in (None,'')): diffs_val.append((a.coordinate,va,vb))
            if (a.has_style or b.has_style) and style_sig(a)!=style_sig(b): diffs_sty.append((a.coordinate,style_sig(a),style_sig(b)))
    res['cells']=n; res['value_diffs']=diffs_val[:15]; res['n_value_diffs']=len(diffs_val); res['n_style_diffs']=len(diffs_sty); res['style_diffs']=diffs_sty[:5]
    res['merges_equal']=sorted(map(str,wo.merged_cells.ranges))==sorted(map(str,wn.merged_cells.ranges))
    wd=lambda ws:{k:round(v.width,2) for k,v in ws.column_dimensions.items() if v.width}
    res['widths_equal']=wd(wo)==wd(wn); res['width_diffs']={k:(wd(wo).get(k),wd(wn).get(k)) for k in set(wd(wo))|set(wd(wn)) if wd(wo).get(k)!=wd(wn).get(k)}
    hd=lambda ws:{k:v.height for k,v in ws.row_dimensions.items() if v.height and (not max_row or k<=max_row)}
    res['heights_equal']=hd(wo)==hd(wn); res['height_diffs']={k:(hd(wo).get(k),hd(wn).get(k)) for k in set(hd(wo))|set(hd(wn)) if hd(wo).get(k)!=hd(wn).get(k)}
    res['freeze']=(wo.freeze_panes,wn.freeze_panes)
    return res
def lo_recalc_check(out_path, sheet, cells):
    """LibreOffice 无界面重算后读取数值"""
    d=tempfile.mkdtemp()
    subprocess.run(['soffice','--headless','--convert-to','xlsx','--outdir',d,out_path],capture_output=True,timeout=120)
    p=os.path.join(d,os.path.basename(out_path))
    wv=openpyxl.load_workbook(p,data_only=True)[sheet]
    wc=openpyxl.load_workbook(out_path,data_only=True)[sheet]  # 文件内缓存值（本系统写入）
    out={c:(wv[c].value, wc[c].value) for c in cells}
    shutil.rmtree(d); return out
if __name__=='__main__':
    print(json.dumps(compare(sys.argv[1],sys.argv[2],sys.argv[3]),ensure_ascii=False,default=str,indent=1))
