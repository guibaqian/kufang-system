"""次日清点（2026-10-19 起）：D 要等 D+1 早上交接表录完才算定；D+1 早上表左边算 D；产品用量 = 早上车上剩余 + 新装 − 退 − 次日早上车上剩余；
待次日清点；手工单；记到哪天 / 在外未交回照旧；新装换人（左右两个配送员）；产品表车牌/配送员跟钢瓶表；下载/打印：没数的人不出现、每人下面一空行、A4 横向一页"""
import os, sys, io, json, re, zipfile, unittest
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from test_all import call, stock, setUpModule, tearDownModule, ROOT, TMP  # noqa
import server, logic, paper, xlsxio

def rows_of(date): return {x['name']: x for x in logic.daily_rows(server.db_c(), date)}
def save(date, ses, form, rows, **kw):
    st, r = call('/api/paper/save', {'data': dict({'date': date, 'session': ses, 'form': form, 'rows': rows}, **kw)}); assert st == 200, r; return r
def moves(pid, date): return sum(r[0] for r in server.db_c().execute("SELECT qty FROM moves WHERE product_id=? AND date=? AND src='daily'", (pid, date)))
def cid(name): return next(k['id'] for k in call('/api/boot')[1]['couriers'] if k['name'] == name)

class T1_NextDay(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        c = server.db_c(); m = logic.daily_map(c); cls.m = m
        for k in ('shb', 'byg', 'xyy'): call('/api/count/save', {'product_id': m[k], 'qty': 500, 'date': '2026-10-01'})
        cls.q0 = {k: stock(m[k])['qty'] for k in ('shb', 'byg', 'xyy', 'c40')}

    def test_1_cyl_and_prod_d_d1(self):
        m = self.m
        # 10.19 早上：左边（退库）是 10.18 的车回来——10.18 在规则开始以前、不动；这部分只进 10.19 库存，不进 10.19 财务表；右边新装
        save('2026-10-19', '早上', 'cyl', [{'rid': 'a', 'plate': '6922', 'driver': '次日甲', 'cells': {'L.e40': '5', 'R.o40': '20'}}])
        save('2026-10-19', '早上', 'prod', [{'rid': 'a', 'plate': '6922', 'cells': {'L.shb': '4', 'L.byg': '2退', 'R.shb': '10', 'R.byg': '5'}}])
        save('2026-10-19', '中午', 'cyl', [{'rid': 'a', 'plate': '6922', 'driver': '次日甲', 'cells': {'L.e40': '8', 'R.o40': '6'}}])
        save('2026-10-19', '中午', 'prod', [{'rid': 'a', 'plate': '6922', 'cells': {'L.shb': '3退', 'R.shb': '2'}, 'hand': [{'product_id': m['xyy'], 'qty': 2}]}])
        self.assertNotIn('次日甲', rows_of('2026-10-18'))
        r = rows_of('2026-10-19')['次日甲']
        self.assertEqual(r['pending'], ['钢瓶', '产品'])                                     # 10.20 早上还没录：待次日清点
        st, x = call('/api/daily/export.xlsx?date=2026-10-19'); sh = xlsxio.read_sheet(x)
        self.assertIn('待次日清点', json.dumps(sh['cells'], ensure_ascii=False)); self.assertNotIn('B5', sh['cells'])   # 下载里不写数
        self.assertFalse(call('/api/daily?date=2026-10-19')[1]['total'].get('o40'))         # 合计不算待清点的
        # 10.20 早上：左边算 10.19
        save('2026-10-20', '早上', 'cyl', [{'rid': 'a', 'plate': '6922', 'driver': '次日甲', 'cells': {'L.e40': '15', 'L.f40': '1', 'R.o40': '18'}}])
        r = rows_of('2026-10-19')['次日甲']; self.assertEqual(r.get('pending'), ['产品'])     # 钢瓶已定，产品表还没录
        save('2026-10-20', '早上', 'prod', [{'rid': 'a', 'plate': '6922', 'cells': {'L.shb': '6', 'L.byg': '1退', 'R.shb': '7'}}])
        r = rows_of('2026-10-19')['次日甲']; v = r['vals']; self.assertNotIn('pending', r)
        # 19 号 = 19 早上出库 + 19 中午入库出库 + 20 早上入库：19 早上钢瓶表左边（18 号的车，旧规则那天）只进库存（se40），不进 19 号财务表
        self.assertEqual((v['o40'], v['e40'], v['f40'], v.get('se40')), (26, 23, 1, 5))         # 早上新装 20 + 中午 6；退 8+15，满 1；早上退 5 只进库存
        self.assertEqual(logic.calc_row(v)['L40'], 2)                                            # 新送 = 26 − 23 − 1
        # 规则第一天（10.19）也用当天早上产品表左边没写退的做起点（1.1.1 当 0 → 出负数，1.1.2 改回）
        self.assertEqual((v['t_shb'], v['i_shb'], v.get('v0_shb'), v['v1_shb'], v['v1_byg']), (12, 3, 4, 6, 1))   # 次日早上左边整块（含写了退的）= 次日车上剩余
        self.assertEqual(logic.calc_row(v)['net_shb'], 7)   # 用量 = 早上车上 4 + 领 12 − 中午退 3 − 次日早上车上 6
        self.assertEqual((v.get('i_byg'), v.get('si_byg'), v.get('v0_byg')), (None, 2, None))   # 10.19 早上写了退的 2：当天入库（只进库存），不从用量里减
        self.assertEqual(logic.calc_row(v)['net_byg'], 4)                                      # 0 + 5 − 次日早上 1
        js = open(os.path.join(ROOT, 'static', 'p_daily.js'), encoding='utf-8').read()   # 负数照样显示，屏幕上标红 + 提示是哪几个数算出来的（打印 / Excel 不标）
        self.assertIn('class="cf neg"', js); self.assertIn("!forPrint && n0(x) < 0", js); self.assertIn('次日早上车上剩余', js)
        self.assertIn('手工单', r['note'])
        r20 = rows_of('2026-10-20')['次日甲']; self.assertEqual(r20['vals'].get('o40'), 18); self.assertFalse(r20['vals'].get('e40')); self.assertEqual(r20['pending'], ['钢瓶', '产品'])
        # 库存 = 新装 − 退（车上剩余不动库存），按日期
        self.assertEqual(moves(m['shb'], '2026-10-19'), -9); self.assertEqual(moves(m['byg'], '2026-10-19'), -3); self.assertEqual(moves(m['byg'], '2026-10-20'), 1)   # 次日早上写退的：10.20 入库; self.assertEqual(moves(m['xyy'], '2026-10-19'), -2)
        self.assertEqual(moves(m['shb'], '2026-10-20'), -7); self.assertEqual(moves(m['c40'], '2026-10-19'), 3)
        self.assertEqual(stock(m['shb'])['qty'], self.q0['shb'] - 16)
        # 下载：领取 T = 新装 12 + 早上车上剩余 4，入库 AI = 退 3 + 次日早上车上剩余 6，当天总出库 AW = 7
        sh = xlsxio.read_sheet(call('/api/daily/export.xlsx?date=2026-10-19')[1])['cells']
        row = next(int(k[1:]) for k, val in sh.items() if k[0] == 'A' and k[1:].isdigit() and val == '次日甲')
        self.assertEqual((sh.get('T%d' % row), sh.get('AI%d' % row), sh.get('AW%d' % row)), (16, 9, 7))

    def test_2_book_date_open_respected(self):
        """1.1.46：人手 book_date 不再生效（记在表的日期 + 次日清点）；open（在外未交回）仍有效"""
        save('2026-10-22', '早上', 'cyl', [{'rid': 'a', 'plate': '6984', 'driver': '次日乙', 'cells': {'L.e40': '4', 'R.o40': '9'}, 'book_date': '2026-10-23'},
                                           {'rid': 'b', 'plate': '9881', 'driver': '次日丙', 'cells': {'L.e40': '2', 'R.o40': '3'}, 'open': True}])
        # 人手记到 23 号被忽略：22 号早上左边在规则下记前一天，右边出库在 22；不出现在 23 的「整行记到」
        self.assertNotIn('次日乙', rows_of('2026-10-23'))
        self.assertNotIn('次日丙', rows_of('2026-10-21')); self.assertNotIn('次日丙', rows_of('2026-10-22'))
        self.assertTrue(any(e['cname'] == '次日丙' for e in call('/api/daily?date=2026-10-22')[1]['open']))

    def test_3_split_row_left_right_couriers(self):
        """新装换人：左边记左边的人，右边记右边的人；次日早上左边算前一天、算左边那个人；车上剩余算左边的人"""
        save('2026-10-25', '中午', 'cyl', [{'rid': 'a', 'plate': '9255', 'driver': '换人丙', 'driver_r': '换人丁', 'cells': {'L.e40': '4', 'R.o40': '9'}}])
        a, b = rows_of('2026-10-25')['换人丙']['vals'], rows_of('2026-10-25')['换人丁']['vals']
        self.assertEqual((a.get('e40'), a.get('o40')), (4, None)); self.assertEqual((b.get('e40'), b.get('o40')), (None, 9))
        save('2026-10-26', '早上', 'cyl', [{'rid': 'a', 'plate': '9255', 'driver': '换人丙', 'driver_r': '换人丁', 'cells': {'L.e40': '3', 'R.o40': '7'}}])
        r = save('2026-10-26', '早上', 'prod', [{'rid': 'a', 'plate': '9255', 'driver': '乱写', 'cells': {'L.shb': '2', 'L.byg': '1退', 'R.shb': '5'}}])
        pr = r['sheet']['rows'][0]; self.assertEqual((pr['driver'], pr['driver_r']), ('换人丙', '换人丁'))   # 产品表跟钢瓶表
        a = rows_of('2026-10-25')['换人丙']['vals']; self.assertEqual((a['e40'], a.get('i_byg'), a['v1_shb'], a['v1_byg']), (7, None, 2, 1))   # 次日早上写了退的也算次日车上剩余，不算前一天的退
        b = rows_of('2026-10-26')['换人丁']['vals']; self.assertEqual((b['o40'], b['t_shb'], b['v0_shb']), (7, 5, 2)); self.assertNotIn('换人丙', rows_of('2026-10-26'))   # 车上剩余跟着车给了右边那位
        # 下一天车给了换人丁：车牌预填右边那位
        self.assertEqual({p['plate']: p for p in call('/api/paper?date=2026-10-27')[1]['plates']}['9255']['name'], '换人丁')

    def test_3b_noon_handover_onboard(self):
        """中午换人：交车时车上剩余（左边没写退）= 左边的人当天收尾剩余、右边的人当天起始剩余"""
        save('2026-10-29', '早上', 'cyl', [{'rid': 'a', 'plate': '5928', 'driver': '交车甲', 'cells': {'R.o40': '1'}}])
        save('2026-10-29', '早上', 'prod', [{'rid': 'a', 'plate': '5928', 'cells': {'R.shb': '3'}}])
        save('2026-10-29', '中午', 'cyl', [{'rid': 'a', 'plate': '5928', 'driver': '交车甲', 'driver_r': '交车乙', 'cells': {'R.o40': '1'}}])
        save('2026-10-29', '中午', 'prod', [{'rid': 'a', 'plate': '5928', 'cells': {'L.shb': '2', 'R.shb': '6'}}])
        save('2026-10-30', '早上', 'cyl', [{'rid': 'a', 'plate': '5928', 'driver': '交车乙', 'cells': {'L.e40': '2'}}])
        save('2026-10-30', '早上', 'prod', [{'rid': 'a', 'plate': '5928', 'cells': {'L.shb': '3'}}])
        r = rows_of('2026-10-29'); a, b = logic.calc_row(r['交车甲']['vals']), logic.calc_row(r['交车乙']['vals'])
        self.assertEqual((a['net_shb'], b['net_shb']), (1, 5))                                 # 甲 3 − 2；乙 2 + 6 − 3；合计 = 领 9 − 车上 3

    def test_4_prod_follows_cyl(self):
        save('2026-10-28', '中午', 'prod', [{'rid': 'p_7539', 'plate': '7539', 'driver': '跟随乙', 'cells': {'R.shb': '1'}}])
        save('2026-10-28', '中午', 'cyl', [{'rid': 'p_7539', 'plate': '7539', 'driver': '跟随甲', 'cells': {'R.o40': '2'}}])
        sh = next(s for s in call('/api/paper?date=2026-10-28')[1]['sheets'] if s['form'] == 'prod')
        self.assertEqual(sh['rows'][0]['driver'], '跟随甲'); self.assertTrue(any(e['k'] == 'driver' and '跟钢瓶页' in e['label'] for e in sh['edits']))
        self.assertEqual(rows_of('2026-10-28')['跟随甲']['vals']['t_shb'], 1); self.assertNotIn('跟随乙', rows_of('2026-10-28'))
        # 在产品表上改配送员不算
        st, r = call('/api/paper/save', {'data': {'id': sh['id'], 'date': '2026-10-28', 'session': '中午', 'form': 'prod', 'rows': [dict(sh['rows'][0], driver='乱改')]}, 'base_ts': sh['updated_at']})
        self.assertEqual(r['sheet']['rows'][0]['driver'], '跟随甲')

    def test_5_downloads_and_print_rows(self):
        """没有数的配送员不出现；每人一行 + 下面一空行；后面没有多余空行；A4 横向、宽高各一页、打印区域；领取人信息不写车牌"""
        D = '2026-11-03'
        call('/api/daily/save', {'data': {'courier_id': cid('6922丁增'), 'out_date': D, 'vals': {'o40': 3}}})
        call('/api/daily/save', {'data': {'courier_id': cid('9255沈兴德'), 'out_date': D, 'vals': {}}})     # 全是 0：不出现
        call('/api/daily/save', {'data': {'courier_id': cid('0668 刘洪江2098'), 'out_date': D, 'vals': {'t_shb': 2}}})
        names = [x['name'] for x in logic.daily_rows(server.db_c(), D)]; self.assertEqual(names, ['6922丁增', '0668 刘洪江2098'])
        x = call('/api/daily/export.xlsx?date=' + D)[1]; z = zipfile.ZipFile(io.BytesIO(x))
        s = z.read('xl/worksheets/sheet1.xml').decode(); wb = z.read('xl/workbook.xml').decode()
        rn = [int(r) for r in re.findall(r'<row r="(\d+)"', s)]; self.assertEqual(max(rn), 4 + 2 * 2)
        c = xlsxio.read_sheet(x)['cells']; self.assertEqual((c.get('A5'), c.get('A7')), ('丁增', '刘洪江2098'))
        self.assertTrue(all(k not in c for k in ('A6', 'A8', 'B6', 'B8', 'BK6', 'BK8')))       # 每人下面空一行
        self.assertIn('orientation="landscape"', s); self.assertIn('paperSize="9"', s); self.assertIn('fitToWidth="1" fitToHeight="1"', s); self.assertIn('fitToPage="1"', s)
        self.assertIn('_xlnm.Print_Area', wb); self.assertIn('$A$1:$BK$8', wb)
        self.assertNotIn('t="shared"', s)                                                       # 共享公式都已改成普通公式（主格可能被删）
        import openpyxl; ws = openpyxl.load_workbook(io.BytesIO(x))['Sheet1']; self.assertEqual(ws['J7'].value, '=D7-H7'); self.assertEqual(ws.max_row, 8)
        # 前端打印同样规则（dHasData + 每人一空行 + A4 横向缩一页）
        js = open(os.path.join(ROOT, 'static', 'p_daily.js'), encoding='utf-8').read()
        self.assertIn('filter(dHasData)', js); self.assertIn("fit: true", js); self.assertIn('class="bl"', js)

    def test_6_old_dates_unchanged(self):
        """规则开始日以前：早上表左边照旧记在表的日期"""
        save('2026-10-18', '早上', 'cyl', [{'rid': 'a', 'plate': '0910', 'driver': '旧规则甲', 'cells': {'L.e40': '6', 'R.o40': '6'}}])
        v = rows_of('2026-10-18')['旧规则甲']['vals']; self.assertEqual((v['e40'], v['o40']), (6, 6)); self.assertNotIn('旧规则甲', rows_of('2026-10-17'))
        self.assertNotIn('pending', rows_of('2026-10-18')['旧规则甲'])
        c = server.db_c(); before = {d: rows_of(d) for d in ('2026-10-18', '2026-10-19', '2026-10-20')}
        paper.migrate_nextday(c); c.commit(); self.assertEqual(paper.migrate_nextday(c), 0)      # 升级只做一次
        self.assertEqual({d: rows_of(d) for d in before}, before)                              # 重算后数不变

    def test_7_version_author(self):
        v = call('/api/version')[1]; self.assertEqual(v['author'], '制作：跳出画框去看画')

    def test_8_nextday_mark_does_not_change_d(self):
        """更正后的产品规则：次日早上写没写「退」都不影响 D；写了退的那部分在次日（实际交回那天）入库"""
        m = self.m; D, N = '2026-11-06', '2026-11-07'
        save(D, '早上', 'cyl', [{'rid': 'a', 'plate': '6720', 'driver': '规则甲', 'cells': {'R.o40': '1'}}])
        save(D, '早上', 'prod', [{'rid': 'a', 'plate': '6720', 'cells': {'L.shb': '4', 'L.byg': '3退', 'R.shb': '10'}}])
        save(D, '中午', 'cyl', [{'rid': 'a', 'plate': '6720', 'driver': '规则甲', 'cells': {'R.o40': '1'}}])
        save(D, '中午', 'prod', [{'rid': 'a', 'plate': '6720', 'cells': {'L.shb': '2退', 'R.shb': '5'}}])
        save(N, '早上', 'cyl', [{'rid': 'a', 'plate': '6720', 'driver': '规则甲', 'cells': {'L.e40': '2'}}])
        res = {}
        for mk in ('6', '6退'):
            save(N, '早上', 'prod', [{'rid': 'a', 'plate': '6720', 'cells': {'L.shb': mk}}]) if mk == '6' else \
                call('/api/paper/save', {'data': {'id': next(x['id'] for x in call('/api/paper?date=' + N)[1]['sheets'] if x['form'] == 'prod'), 'date': N, 'session': '早上', 'form': 'prod',
                                                  'rows': [{'rid': 'a', 'plate': '6720', 'cells': {'L.shb': mk}}]}})
            r = rows_of(D)['规则甲']; v = r['vals']; self.assertNotIn('pending', r)
            res[mk] = (v.get('v0_shb'), v.get('t_shb'), v.get('i_shb'), v.get('v1_shb'), logic.calc_row(v)['net_shb'], v.get('v0_byg'), moves(m['shb'], D), moves(m['shb'], N))
        self.assertEqual(res['6'][:6], (4, 15, 2, 6, 11, None))       # 4 + 15 − 2 − 6 = 11；早上写了退的 byg 不算 D 的起始
        self.assertEqual(res['6退'][:6], res['6'][:6])                  # 次日早上写没写退，D 一样
        self.assertEqual((res['6'][6], res['6'][7], res['6退'][7]), (-13, 0, 6))   # 库存：D 扣 新装 − 中午退；写了退的 6 在 N 入库
        self.assertEqual(moves(m['byg'], D), 3)                         # D 早上写退的 3：D 入库
        self.assertEqual(rows_of('2026-11-05')['规则甲']['vals'].get('v1_byg'), 3)   # 财务上算前一天的「次日车上剩余」

    def test_9_long_note_goes_to_blank_row(self):
        """备注太长（超过约 14 个汉字）：下载 / 打印放到这人下面那一空行（整行合并、靠左换行），备注格写「见下行」；短备注照旧"""
        D = '2026-11-15'; long = '下午在第二人民医院补送 2 瓶 40L（单号 A1023），空瓶 3 个明早交回，财务核对后再入账'
        call('/api/daily/save', {'data': {'courier_id': cid('6922丁增'), 'out_date': D, 'vals': {'o40': 3}, 'note': long}})
        call('/api/daily/save', {'data': {'courier_id': cid('0668 刘洪江2098'), 'out_date': D, 'vals': {'o40': 2}, 'note': '少交 1 个空瓶'}})
        x = call('/api/daily/export.xlsx?date=' + D)[1]; z = zipfile.ZipFile(io.BytesIO(x)); s = z.read('xl/worksheets/sheet1.xml').decode()
        c = xlsxio.read_sheet(x)['cells']
        self.assertEqual((c.get('BK5'), c.get('A6'), c.get('BK7'), c.get('A8')), ('见下行', long, '少交 1 个空瓶', None))
        self.assertIn('<mergeCell ref="A6:BK6"/>', s); self.assertNotIn('A8:BK8', s); self.assertEqual(max(int(r) for r in re.findall(r'<row r="(\d+)"', s)), 8)
        import openpyxl; ws = openpyxl.load_workbook(io.BytesIO(x))['Sheet1']
        self.assertEqual((ws['A6'].alignment.horizontal, ws['A6'].alignment.wrap_text), ('left', True)); self.assertEqual(ws.page_setup.orientation, 'landscape')
        self.assertTrue(xlsxio.note_long(long)); self.assertFalse(xlsxio.note_long('少交 1 个空瓶，明早补'))
        js = open(os.path.join(ROOT, 'static', 'p_daily.js'), encoding='utf-8').read(); self.assertIn("noteLong = t => noteW(t) > 28", js); self.assertIn('见下行', js)
        # 打印 = 下载的 Excel 原样转 HTML：原表列宽、灰底（更换 / 实际配送）、蓝底（新送）、宋体字号、行高；长备注整行合并
        h = call('/api/daily/print?date=' + D)[1]['html']
        self.assertIn('2026年11月15号产品出入库情况', h); self.assertIn('font-size:24pt', h); self.assertIn('height:39pt', h); self.assertIn('height:27pt', h)
        self.assertIn('background:#a6a6a6', h); self.assertIn('background:#91ace0', h); self.assertNotIn('<col style="width:122px">', h); self.assertIn('height:15pt', h)
        self.assertIn('colspan="34"', h); self.assertIn('>见下行<', h); self.assertIn(long[:10], h); self.assertNotIn('>0<', h)
        self.assertIn("/api/daily/print?date=", js)

    def test_9b_gap_combined_period(self):
        """第二天早上没清点到这个人：前一天接着「待清点」；到下次早上清点时几天合起来算（最后一天显示「1–2号合计」，中间那天「并入 2号合计」）"""
        A, B, C = '2026-12-01', '2026-12-02', '2026-12-03'
        save(A, '早上', 'cyl', [{'rid': 'a', 'plate': '1664c', 'driver': '断档甲', 'cells': {'R.o40': '10'}}])
        save(A, '早上', 'prod', [{'rid': 'a', 'plate': '1664c', 'cells': {'L.shb': '3', 'R.shb': '10'}}])
        save(B, '早上', 'cyl', [{'rid': 'b', 'plate': '0928', 'driver': '断档乙', 'cells': {'R.o40': '1'}}])   # 2 号早上的表里没有断档甲
        save(B, '早上', 'prod', [{'rid': 'b', 'plate': '0928', 'cells': {'R.shb': '1'}}])
        r = rows_of(A)['断档甲']; self.assertTrue(r['pending'] and r['gap'])                       # 待清点（不是待次日清点）
        save(B, '中午', 'cyl', [{'rid': 'a', 'plate': '1664c', 'driver': '断档甲', 'cells': {'R.o40': '5'}}])
        save(B, '中午', 'prod', [{'rid': 'a', 'plate': '1664c', 'cells': {'L.shb': '1退', 'R.shb': '4'}}])
        self.assertTrue(rows_of(B)['断档甲']['pending'])
        save(C, '早上', 'cyl', [{'rid': 'a', 'plate': '1664c', 'driver': '断档甲', 'cells': {'L.e40': '12'}}])
        save(C, '早上', 'prod', [{'rid': 'a', 'plate': '1664c', 'cells': {'L.shb': '2'}}])
        a, b = rows_of(A)['断档甲'], rows_of(B)['断档甲']
        self.assertNotIn('pending', a); self.assertEqual(a['merged_into'], B)
        self.assertEqual(b['period'], [A, B]); self.assertTrue(b['note'].startswith('1–2号合计'))
        v = b['vals']; self.assertEqual((v['o40'], v['e40']), (15, 12))
        self.assertEqual((v['v0_shb'], v['t_shb'], v['i_shb'], v['v1_shb'], logic.calc_row(v)['net_shb']), (3, 14, 1, 2, 14))   # 3 + 14 − 1 − 2
        self.assertFalse(call('/api/daily?date=' + A)[1]['total'].get('o40'))                      # 并入的那天不算合计
        self.assertEqual(call('/api/daily?date=' + B)[1]['total'].get('o40'), 15 + 1 if '断档乙' in rows_of(B) and not rows_of(B)['断档乙'].get('pending') else 15)
        ca = xlsxio.read_sheet(call('/api/daily/export.xlsx?date=' + A)[1])['cells']; cb = xlsxio.read_sheet(call('/api/daily/export.xlsx?date=' + B)[1])['cells']
        ra = next(k[1:] for k, x in ca.items() if k[0] == 'A' and x == '断档甲'); rb = next(k[1:] for k, x in cb.items() if k[0] == 'A' and x == '断档甲')
        self.assertEqual((ca.get('BK' + ra), ca.get('B' + ra)), ('并入2号合计', None)); self.assertEqual(cb.get('B' + rb), 15); self.assertTrue(cb.get('BK' + rb).startswith('1–2号合计'))
        self.assertEqual(moves(self.m['shb'], A) + moves(self.m["shb"], B), -14)                   # 库存照常按天：新装 14 − 中午退 1，加断档乙领的 1

    def test_9c_zzq_style_gap_19_21_no_manual_book(self):
        """跨日配送类：19 出库、20 早上没此人、21 早上清到 → 19 并入、20 显示「19–20号合计」；不要靠人手 book_date"""
        A, B, C = '2026-12-10', '2026-12-11', '2026-12-12'
        save(A, '早上', 'cyl', [{'rid': 'a', 'plate': '0928z', 'driver': '合算甲', 'cells': {'R.o40': '40', 'R.o10': '9'}}])
        save(A, '早上', 'prod', [{'rid': 'a', 'plate': '0928z', 'cells': {'L.shb': '17', 'L.byg': '10'}}])
        # 20 早上表有别人、没有合算甲
        save(B, '早上', 'cyl', [{'rid': 'b', 'plate': '1664c', 'driver': '合算乙', 'cells': {'R.o40': '1'}}])
        save(B, '早上', 'prod', [{'rid': 'b', 'plate': '1664c', 'cells': {'R.shb': '1'}}])
        ra = rows_of(A)['合算甲']; self.assertTrue(ra.get('pending') and ra.get('gap'))
        # 21 早上清到
        save(C, '早上', 'cyl', [{'rid': 'a', 'plate': '0928z', 'driver': '合算甲', 'cells': {'L.e40': '39', 'L.e10': '3'}}])
        save(C, '早上', 'prod', [{'rid': 'a', 'plate': '0928z', 'cells': {'L.shb': '16', 'L.byg': '5'}}])
        a, b = rows_of(A)['合算甲'], rows_of(B)['合算甲']
        self.assertNotIn('pending', a); self.assertEqual(a['merged_into'], B)
        self.assertEqual(b['period'], [A, B]); self.assertTrue(b['note'].startswith('10–11号合计'))
        v = b['vals']; self.assertEqual((v['o40'], v['e40']), (40, 39))
        self.assertEqual((v.get('v0_shb'), v.get('v1_shb'), logic.calc_row(v)['net_shb']), (17, 16, 1))
        self.assertNotIn('合算甲', rows_of(C))   # 清点日早上只交回、当天没出车：不进当天汇总冒用量
        # 前端已去掉「记到哪天」日期选项
        js = open(os.path.join(ROOT, 'static', 'p_daily.js'), encoding='utf-8').read()
        self.assertIn('在外未交回…', js); self.assertNotIn('记到哪天 / 在外未交回', js)
        self.assertIn('lyNd', js); self.assertIn('规则开始日', js)

    def test_9d_clean_rows_strips_manual_book_on_ui_save(self):
        """保存交接表时前端传 book_date=null：残留「记到哪天」被清掉，改走次日清点"""
        D, N = '2026-12-20', '2026-12-21'
        save(D, '早上', 'cyl', [{'rid': 'a', 'plate': '7539z', 'driver': '清记甲', 'cells': {'R.o40': '8'}, 'book_date': N}])
        # 上一行若仍吃 book_date，会把出库记到 N；1.1.46 clean 后应记在 D 且待次日清点
        r = rows_of(D).get('清记甲') or {}
        self.assertTrue(r.get('pending'), r)
        self.assertEqual((r.get('vals') or {}).get('o40'), 8)
        self.assertNotIn('清记甲', rows_of(N))

if __name__ == '__main__': unittest.main()
