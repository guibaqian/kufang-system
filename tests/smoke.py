"""全新安装冒烟测试：python tests/smoke.py http://127.0.0.1:8095 （服务器用空数据目录启动）
接口：每个模块 新建→修改→删除→回收站→恢复；密码；打印数据；所有下载（与原模板对比）；备份/导出/恢复。
界面：桌面每页、手机每页无报错；交接表录入/修改/修改记录；表格设置；手机交接表；各页「打印」与「⋯」菜单。"""
import sys, os, json, io, zipfile, re, tempfile, urllib.request, urllib.error, asyncio
U = sys.argv[1] if len(sys.argv) > 1 else 'http://127.0.0.1:8095'
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'app')); sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import compare
OK, BAD = [], []
def check(name, cond, info=''):
    (OK if cond else BAD).append(name + ('' if cond else '  ← ' + str(info)[:300]))
def call(p, b=None, pin='1002', raw=None):
    if b is None and raw is None: rq = urllib.request.Request(U + p)
    elif raw is not None: rq = urllib.request.Request(U + p, raw, {'Content-Type': 'application/octet-stream', 'X-Pin': pin})
    else: rq = urllib.request.Request(U + p, json.dumps(dict(b, pin=pin)).encode(), {'Content-Type': 'application/json'})
    try:
        with urllib.request.urlopen(rq, timeout=60) as r:
            d = r.read(); return r.status, (json.loads(d) if 'json' in r.headers.get('Content-Type', '') else d)
    except urllib.error.HTTPError as e:
        t = e.read()
        try: return e.code, json.loads(t or b'{}')
        except Exception: return e.code, t

def crud(name, save, data, edit, delete_path, ent, read=None):
    st, r = call(save, {'data': data}); check(name + ' 新建', st == 200, r)
    if st != 200: return None
    i = r['id'] if 'id' in r else r['ids'][0]
    if edit:
        st, r2 = call(save, {'data': dict(data, id=i, **edit)}); check(name + ' 修改', st == 200, r2)
    st, im = call('/api/impact?entity=%s&id=%d' % (ent, i)); check(name + ' 删除前影响清单', st == 200 and 'effects' in im, im)
    if ent in ('out', 'arrival', 'loan', 'daily', 'paper', 'slip'):
        st, rv = call('/api/record?entity=%s&id=%d' % (ent, i)); check(name + ' 记录来源查看', st == 200 and rv.get('id') == i, rv)
    st, r3 = call(delete_path, {'id': i}); check(name + ' 删除', st == 200, r3)
    tr = call('/api/trash')[1]['items']; check(name + ' 在回收站', any(x['entity'] == ent and x['id'] == i for x in tr), ent)
    st, r4 = call('/api/restore', {'entity': ent, 'id': i}); check(name + ' 恢复', st == 200, r4)
    return i

def api_part():
    st, v = call('/api/version'); check('版本号接口', st == 200 and v.get('version') == '1.1.2' and v.get('changelog'), v)
    st, r = call('/api/settings/save', {'data': {'ui_names': {'出库': '发货'}, 'print': {'company': '测试公司'}}}); check('名称设置保存', st == 200, r)
    b = call('/api/boot')[1]; check('名称设置生效', b['settings']['ui_names'] == {'出库': '发货'} and b['settings']['print']['footer'] is True, b['settings'].get('ui_names'))
    call('/api/settings/save', {'data': {'ui_names': {}, 'print': {'company': ''}}})
    tpl = os.path.join(ROOT, 'templates', '产品出入库情况_原表.xlsx')
    st, x = call('/api/daily/export.xlsx?date=2026-10-17')
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f: f.write(x)
    import openpyxl; mr = openpyxl.load_workbook(f.name)['Sheet1'].max_row
    r = compare.compare(tpl, 'Sheet1', f.name, max_row=mr)   # 后面多余的空行已去掉（用户要求），只比有内容的行
    vd = [d for d in r.get('value_diffs', []) if not (re.match(r'^(A|AH)\d+$', d[0]) and isinstance(d[1], str) and isinstance(d[2], str) and d[1].endswith(d[2]) and d[1] != d[2])]   # 领取人信息不写车牌（6922丁增 → 丁增，用户要求）
    check('产品出入库情况 10.17 下载 = 原表（领取人信息不带车牌）', len(vd) == 0 and r['n_value_diffs'] <= 20 and r['n_style_diffs'] == 0 and r['merges_equal'] and set(r['width_diffs']) <= {'A', 'S', 'AW'}, (vd[:3], r['width_diffs'], r['n_value_diffs'], r['n_style_diffs'], r['style_diffs'][:2]))   # 列宽：A 按名字收窄、数字列等宽（1.1.1 用户要的排版）
    import reports
    for k, cfg in reports.TABS.items():
        if k == 'yard4': continue
        st, x = call('/api/report.xlsx?sums=0&key=' + k)
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f: f.write(x)
        r = compare.compare(os.path.join(ROOT, 'templates', reports.BOOKS[cfg['book']]), cfg['sheet'], f.name)
        check('报表下载 = 原模板 ' + k, r['n_value_diffs'] == 0 and r['n_style_diffs'] == 0 and r['merges_equal'], r.get('value_diffs', [])[:3])
    st, x = call('/api/report.xlsx?key=lox'); import xlgrid; sh = xlgrid.read_sheet(x, '液氧表'); ss = reports.auto_sums('lox', call('/api/report?key=lox')[1])
    check('液氧表下载含自动月合计', all(reports._v(sh, 'A', s['after'] + 1 + i) == '合计' for i, s in enumerate(ss)) and len(ss) >= 1, ss)
    st, e = call('/api/whoami', {}, pin='9999'); check('错误密码被拒', st in (401, 403), st)
    st, e = call('/api/product/save', {'data': {'name': 'x'}}, pin='0000'); check('错误密码不能保存', st in (401, 403), st)
    boot = call('/api/boot')[1]; ps = {p['name']: p['id'] for p in boot['products']}
    check('首次启动导入了配送员', len(boot['couriers']) >= 10, len(boot['couriers']))
    pid = crud('产品', '/api/product/save', {'name': '冒烟产品', 'unit': '个', 'category': '其他', 'remark': '备注A'}, {'spec': 'S1'}, '/api/product/delete', 'product')
    st, _ = call('/api/count/save', {'product_id': pid, 'qty': 100, 'date': '2026-10-01', 'remark': '盘点'}); check('盘点', st == 200)
    crud('到货', '/api/arrival/save', {'date': '2026-10-02', 'product_id': pid, 'qty': 6, 'u8_status': '待入U8', 'remark': 'r'}, {'qty': 7}, '/api/arrival/delete', 'arrival')
    crud('出库', '/api/out/save', {'date': '2026-10-03', 'kind': '自提出库', 'customer': '冒烟医院', 'source': 'app手工单', 'remark': 'r', 'items': [{'product_id': pid, 'qty': 2}]},
         {'remark': '改了'}, '/api/out/delete', 'out')
    oid = call('/api/out/save', {'data': {'date': '2026-10-03', 'kind': '领用出库', 'recipient': '张三', 'purpose': '测', 'items': [{'product_id': pid, 'qty': 1}]}})[1]['id']
    lid = crud('借用', '/api/loan/save', {'date': '2026-10-04', 'handler': '张三', 'user_to': '李医生', 'orgs': ['一院'], 'need_return': 1, 'items': [{'product_id': pid, 'qty': 3, 'org': '一院'}]},
               {'remark': '改'}, '/api/loan/delete', 'loan')
    one = call('/api/loan/one?id=%d' % lid)[1]
    st, rr = call('/api/loan/return', {'loan_id': lid, 'date': '2026-10-05', 'remark': '还一个', 'items': [{'loan_item_id': one['items'][0]['id'], 'qty': 1}]}); check('归还', st == 200, rr)
    crud('库房杂物', '/api/misc/save', {'name': '冒烟杂物', 'category': '工程', 'qty': '一堆', 'remark': 'r'}, {'location': '门口'}, '/api/misc/action', 'temp') if False else None
    st, m = call('/api/misc/save', {'data': {'name': '冒烟杂物', 'category': '工程', 'qty': '一堆', 'remark': 'r'}}); check('库房杂物 新建', st == 200, m)
    st, _ = call('/api/misc/save', {'data': {'id': m['id'], 'name': '冒烟杂物2', 'category': '工程'}}); check('库房杂物 修改', st == 200)
    st, pe = call('/api/pending/save', {'data': {'taker': '张总', 'what': '冒烟制氧机', 'qty': '1', 'remark': 'r'}}); check('待补全 新建', st == 200, pe)
    st, _ = call('/api/pending/close', {'id': pe['id'], 'note': '已补'}); check('待补全 补全', st == 200, _)
    sl = crud('称重单', '/api/slip/save', {'plate': '藏A冒烟', 'gross': 30000, 'tare': 10000, 'note': 'r'}, {'gross': 31000}, '/api/slip/delete', 'slip')
    crud('配送员', '/api/courier/save', {'name': '冒烟配送员'}, {'kind': '配送员'}, '/api/courier/delete', 'courier')
    # 交接表：新建 → 对账后修改 → 修改记录 → 删除 → 回收站恢复
    st, r = call('/api/paper/save', {'data': {'date': '2026-10-06', 'session': '早上', 'form': 'cyl', 'rows': [{'rid': 'a', 'plate': '6922', 'driver': '冒烟司机', 'cells': {'R.o40': '10+2', 'L.e40': '9'}}]}})
    check('交接表 新建', st == 200, r)
    rows = r['sheet']['rows']; rows[0]['cells']['R.o40'] = '13'
    st, r2 = call('/api/paper/save', {'data': {'id': r['id'], 'date': '2026-10-06', 'session': '早上', 'form': 'cyl', 'rows': rows}, 'base_ts': r['sheet']['updated_at']}, pin='1001')
    check('交接表 对账后修改', st == 200 and r2['sheet']['edits'][-1]['old'] == '10+2' and r2['sheet']['edits'][-1]['new'] == '13', r2)
    d = call('/api/daily?date=2026-10-06')[1]; check('交接表 → 汇总', any(m['name'] == '冒烟司机' and m['vals'].get('o40') == 13 for m in d['merged']), d['merged'])
    st, _ = call('/api/paper/delete', {'id': r['id']}); check('交接表 删除', st == 200)
    check('交接表 在回收站', any(x['entity'] == 'paper' and x['id'] == r['id'] for x in call('/api/trash')[1]['items']))
    call('/api/restore', {'entity': 'paper', 'id': r['id']}); d = call('/api/daily?date=2026-10-06')[1]
    check('交接表 恢复后汇总回来', any(m['name'] == '冒烟司机' for m in d['merged']))
    st, _ = call('/api/paper/row', {'data': {'date': '2026-10-06', 'session': '中午', 'form': 'prod', 'row': {'plate': '9255', 'driver': '冒烟司机', 'cells': {'R.shb': '3', 'L.shb': '1退'}}}}); check('交接表 手机一行', st == 200, _)
    lay = call('/api/paper/layout')[1]; st, _ = call('/api/paper/layout', {'layout': lay['layout'], 'plates': lay['plates'], 'sessions': lay['sessions']}, pin='1001'); check('表格设置 保存', st == 200, _)
    # 报表：保存一行
    t = call('/api/report?key=yard4')[1]; rws = t['rows'] + [{'c': {'A': {'v': '2026.10.20'}, 'B': {'v': '冒烟'}}}]
    st, _ = call('/api/report/save', {'tabs': {'yard4': {'rows': rws, 'merges': t['merges'], 'version': t['version']}}, 'summary': '冒烟'}); check('报表 保存', st == 200, _)
    st, _ = call('/api/report/save', {'tabs': {'yard4': {'rows': rws, 'merges': t['merges'], 'version': t['version']}}, 'summary': '冒烟'}); check('报表 版本冲突 409', st == 409, st)
    # 打印数据 + 下载
    for path in ('/api/count/sheet.xlsx', '/api/pending/list.xlsx', '/api/out/export.xlsx', '/api/slips.xlsx', '/api/out/lingyong?xlsx=1', '/api/loans.xlsx', '/api/loan/org.xlsx?org=' + urllib.parse.quote('一院'),
                 '/api/batch/trace.xlsx?product_id=%d&batch_no=' % pid):
        st, j = call(path + ('&' if '?' in path else '?') + 'fmt=json'); check('打印数据 ' + path, st == 200 and isinstance(j, dict) and j.get('sheets'), (st, str(j)[:120]))
        st, x = call(path); ok = st == 200 and x[:2] == b'PK'
        if ok:
            import openpyxl; openpyxl.load_workbook(io.BytesIO(x))
        check('下载 ' + path, ok, st)
    for path in ('/api/loan/slip.xlsx?id=%d' % lid, '/api/out/slip.xlsx?id=%d' % oid):
        st, x = call(path); check('下载 ' + path, st == 200 and x[:2] == b'PK', st)
    st, x = call('/api/slip.docx?id=%d' % sl); txt = ''
    if st == 200: txt = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', zipfile.ZipFile(io.BytesIO(x)).read('word/document.xml').decode()))
    check('称重单 Word', '藏A冒烟' in txt, st)
    # 备份 / 导出 / 恢复
    st, b = call('/api/backup/now', {}, pin='1001'); check('立即备份', st == 200, b)
    bl = call('/api/backup/list')[1]; check('备份列表', bl.get('items'), bl)
    st, z = call('/api/export.zip'); check('导出全部', st == 200 and z[:2] == b'PK', st)
    n0 = len(call('/api/products')[1]['items'])
    st, r = call('/api/import', raw=z, pin='1001'); check('导入/恢复', st == 200, r)
    check('恢复后数据一致', len(call('/api/products')[1]['items']) == n0)
    for p in ['boot', 'home', 'products', 'out', 'daily', 'arrivals', 'batches', 'pending', 'misc', 'people', 'u8', 'audit', 'trash', 'backup/list', 'counts', 'loans', 'slips', 'reports',
              'paper', 'paper/layout', 'out/lingyong', 'loan/outstanding', 'customers', 'daily/dates', 'inbox?token=abc']:
        st, _ = call('/api/' + p); check('GET /api/' + p, st == 200, st)

DESK = ['#home', '#out/single', '#out/batch', '#out/list', '#out/ly', '#daily/cyl', '#daily/prod', '#daily/sum', '#daily/layout', '#arrival', '#stock/list', '#stock/count', '#stock/batch', '#stock/u8',
        '#misc', '#pending', '#loan/new', '#loan/list', '#loan/org', '#loan/batch', '#report/conv', '#slip', '#backup', '#backup/trash', '#backup/audit', '#settings'] + ['#settings/' + k for k in ('about', 'names', 'defaults', 'print', 'reportcols', 'layout', 'plates', 'daily', 'ocr', 'options', 'couriers', 'products', 'people', 'orgs', 'autostart', 'backup', 'phone', 'other')]
PHONE = ['#', '#out', '#cyl', '#cyl/prod:6922', '#misc', '#pend', '#mine', '#loan', '#cam/abc', '#settings']
async def ui_part():
    from playwright.async_api import async_playwright
    async with async_playwright() as p:
        b = await p.chromium.launch(executable_path='/usr/bin/google-chrome', args=['--no-sandbox'])
        for kind, items, vp, path in (('桌面', DESK, {'width': 1440, 'height': 900}, '/'), ('手机', PHONE, {'width': 390, 'height': 844}, '/m/')):
            ctx = await b.new_context(viewport=vp, is_mobile=kind == '手机'); pg = await ctx.new_page(); errs = []
            pg.on('pageerror', lambda e: errs.append(str(e))); pg.on('console', lambda m: m.type == 'error' and errs.append(m.text))
            pg.on('dialog', lambda d: asyncio.ensure_future(d.accept()))
            await pg.goto(U + path); await pg.wait_for_timeout(800)
            for h in items:
                n0 = len(errs); await pg.evaluate('h => location.hash = h', h); await pg.wait_for_timeout(900)
                txt = await pg.evaluate('document.body.innerText')
                check('%s 页面 %s' % (kind, h), len(errs) == n0 and '出错' not in txt[:200], errs[n0:])
                if kind == '桌面':   # 打开第一行的「⋯」菜单
                    dots = await pg.query_selector('table .ra-dots')
                    if dots:
                        await dots.click(); await pg.wait_for_timeout(150)
                        check('%s ⋯菜单 %s' % (kind, h), await pg.query_selector('.ra-more.on .ra-menu a') is not None)
                        await pg.mouse.click(5, 5)
            if kind == '桌面':
                # 交接表界面录入 → 保存（密码）→ 改一格 → 修改记录
                await pg.evaluate("location.hash='#daily/cyl'"); await pg.wait_for_timeout(900)
                await pg.click('#ppSes button:has-text("晚上")'); await pg.wait_for_timeout(300)
                await pg.click('input[data-r="0"][data-k="R.o40"]'); await pg.keyboard.type('5+1')
                await pg.click('#ppSave'); await pg.wait_for_timeout(400)
                if await pg.query_selector('.pin-in'): await pg.fill('.pin-in', '1002'); await pg.keyboard.press('Enter')
                await pg.wait_for_timeout(1200)
                await pg.click('#ppSes button:has-text("晚上")'); await pg.wait_for_timeout(300)
                await pg.fill('input[data-r="0"][data-k="R.o40"]', '7'); await pg.dispatch_event('input[data-r="0"][data-k="R.o40"]', 'input')
                await pg.click('#ppSave'); await pg.wait_for_timeout(400)
                if await pg.query_selector('.pin-in'): await pg.fill('.pin-in', '1002'); await pg.keyboard.press('Enter')
                await pg.wait_for_timeout(1200); await pg.click('#ppSes button:has-text("晚上")'); await pg.wait_for_timeout(300)
                check('桌面 交接表 改过的格子有标记', await pg.query_selector('td.ed input[data-k="R.o40"]') is not None)
                h = await pg.query_selector('#ppHist')
                if h: await h.click(); await pg.wait_for_timeout(300)
                check('桌面 交接表 修改记录', h is not None and '5+1' in await pg.evaluate("document.querySelector('.dlg') ? document.querySelector('.dlg').innerText : ''"))
                # 打印按钮
                for hh, sel in (('#daily/sum', '#dPrint'), ('#report/conv', '#rPrint'), ('#out/list', 'a:has-text("打印")')):
                    await pg.evaluate('h => location.hash = h', hh); await pg.wait_for_timeout(900); n0 = len(errs)
                    el = await pg.query_selector(sel)
                    if el:
                        try: await el.click(timeout=2000)
                        except Exception: pass
                        await pg.wait_for_timeout(500)
                    check('桌面 打印 %s' % hh, len(errs) == n0, errs[n0:])
            else:
                await pg.evaluate("location.hash='#cyl/cyl:6922'"); await pg.wait_for_timeout(900)
                await pg.click('#pS button:has-text("中午")'); await pg.wait_for_timeout(300)
                await pg.fill('.ln.pp[data-k="R.o10"] input', '2'); await pg.click('#pSave'); await pg.wait_for_timeout(500)
                for k in '1002': await pg.click('.pinpad [data-k="%s"]' % k)
                await pg.wait_for_timeout(1200)
                check('手机 交接表 保存', '已记入' in await pg.evaluate('document.body.innerText'))
            await ctx.close()
        await b.close()

api_part()
asyncio.run(ui_part())
print('通过 %d 项，失败 %d 项' % (len(OK), len(BAD)))
for x in BAD: print('  ✗', x)
sys.exit(1 if BAD else 0)
